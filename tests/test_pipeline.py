import os
import sys

import cv2
import numpy as np
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tests.synth import make_sheet  # noqa: E402
from traceview.engine import Engine  # noqa: E402
from traceview.excel_out import PendingEditsError, export_xlsx  # noqa: E402
from traceview.project import SourceMismatch, load_project, save_project  # noqa: E402
from traceview.svg_out import export_svg  # noqa: E402


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    rgb, truth = make_sheet()
    p = tmp_path_factory.mktemp("img") / "sheet.png"
    cv2.imwrite(str(p), cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
    e = Engine()
    e.open(str(p))
    e.analyze()
    return e, truth, p


def _recall(band, peaks, tol=6):
    hit = 0
    for x, y, kind in peaks:
        ok = any(v.kind == kind and abs(v.x - x) <= tol and abs(v.y - y) <= tol for v in band.vertices)
        hit += ok
    return hit / len(peaks)


def test_three_bands_found_with_baselines(sheet):
    e, truth, _ = sheet
    assert len(e.bands) == 3
    for b, t in zip(e.bands, truth.values()):
        assert abs(b.y0 - t["baseline"]) < 6


def test_each_band_traced_to_end(sheet):
    e, _, _ = sheet
    for b in e.bands:
        assert b.complete, b.name
        assert b.end_x > 2200


def test_tips_are_reached(sheet):
    """TR-02: every peak / valley tip is reached (also tips of narrow, retraced spikes)."""
    e, truth, _ = sheet
    for b, t in zip(e.bands, truth.values()):
        assert _recall(b, t["peaks"]) >= 0.95, (b.name, _recall(b, t["peaks"]))


def test_path_stays_on_own_band(sheet):
    e, truth, _ = sheet
    for b, t in zip(e.bands, truth.values()):
        assert np.abs(b.path[:, 1] - t["baseline"]).max() < 175


def test_measurement_matches_independent_calculation(sheet):
    e, truth, _ = sheet
    for b, t in zip(e.bands, truth.values()):
        m = e.measures(b)
        top = min(p[1] for p in t["peaks"])
        bot = max(p[1] for p in t["peaks"])
        assert abs(m["max_value"] - (b.y0 - top) * 1.0) < 6
        assert abs(m["min_value"] - (b.y0 - bot) * 1.0) < 6
        assert abs(m["span"] - (bot - top)) < 8
        assert m["up_amp"] >= 0 and m["down_amp"] >= 0


def test_undetermined_until_confirmed(sheet):
    e, _, _ = sheet
    b = e.bands[0]
    b.confirmed = False
    assert not e.is_determinate(b)
    b.confirmed = True
    assert e.is_determinate(b)
    b.confirmed = False


def test_move_vertex_retraces_and_undo(sheet):
    e, _, _ = sheet
    b = e.bands[1]
    k = len(b.vertices) // 2
    v = b.vertices[k]
    n_before = len(b.path)
    e.set_pending_move(1, k, v.x + 2, v.y)
    assert e.has_pending()
    ok, msg = e.apply_pending(1)
    assert ok, msg
    assert not e.has_pending()
    assert any(p.source == "manual" for p in b.vertices)
    assert b.complete                          # later part still traced
    assert e.undo(1)
    assert len(b.path) == n_before


def test_failed_edit_keeps_previous_result(sheet):
    e, _, _ = sheet
    b = e.bands[2]
    before = b.path.copy()
    e.set_pending_move(2, 3, 5, 5)             # nowhere near a line
    ok, msg = e.apply_pending(2)
    assert not ok and msg
    assert np.array_equal(b.path, before)
    e.cancel_pending(2)


def test_export_blocked_while_pending(sheet, tmp_path):
    e, _, _ = sheet
    v = e.bands[0].vertices[5]
    e.set_pending_move(0, 5, v.x + 1, v.y)
    with pytest.raises(PendingEditsError):
        export_xlsx(e, str(tmp_path / "x.xlsx"))
    with pytest.raises(PendingEditsError):
        export_svg(e, str(tmp_path / "x.svg"))
    e.cancel_pending(0)


def test_excel_and_svg(sheet, tmp_path):
    from openpyxl import load_workbook
    e, _, _ = sheet
    out = export_xlsx(e, str(tmp_path / "m.xlsx"))
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["測定結果と画像", "復元座標", "頂点と折返し"]
    assert len(wb["測定結果と画像"]._images) == 2
    assert wb["復元座標"].max_row > 100
    assert wb["頂点と折返し"].max_row > 100
    svg = export_svg(e, str(tmp_path / "t.svg"))
    assert open(svg, encoding="utf-8").read().count("<path") >= 3


def test_project_roundtrip_and_source_check(sheet, tmp_path):
    e, _, p = sheet
    e.bands[0].confirmed = True
    e.calibration.mm_per_px_y = 0.05
    f = str(tmp_path / "w.trproj")
    save_project(e, f)
    e2 = Engine()
    load_project(e2, f)
    assert len(e2.bands) == 3
    assert np.array_equal(e2.bands[1].path, e.bands[1].path)
    assert e2.bands[0].confirmed and e2.calibration.mm_per_px_y == 0.05
    assert len(e2.bands[1].vertices) == len(e.bands[1].vertices)
    # a different original must be refused
    other, _ = make_sheet(seed=7)
    q = tmp_path / "other.png"
    cv2.imwrite(str(q), cv2.cvtColor(other, cv2.COLOR_RGB2BGR))
    with pytest.raises(SourceMismatch):
        load_project(Engine(), f, source_path=str(q))
    e.bands[0].confirmed = False
    e.calibration.mm_per_px_y = 0.0


def test_original_file_untouched(sheet):
    import hashlib
    e, _, p = sheet
    assert hashlib.sha256(open(p, "rb").read()).hexdigest() == e.digest


def test_sample_sheet_if_present():
    sample = os.path.join(os.path.dirname(__file__), "..", "samples", "sample_B1-01.pdf")
    if not os.path.exists(sample):
        pytest.skip("sample not available")
    e = Engine()
    e.open(sample)
    e.analyze()
    assert len(e.bands) == 3
    assert all(len(b.path) > 1000 for b in e.bands)
    assert all(len(b.vertices) > 100 for b in e.bands)
