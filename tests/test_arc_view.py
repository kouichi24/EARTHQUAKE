import os
import sys

import cv2
import numpy as np
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tests.synth import make_sheet  # noqa: E402
from traceview.engine import Engine  # noqa: E402
from traceview.excel_out import export_xlsx  # noqa: E402
from traceview.measure import Calibration, arc_active, corrected_path, corrected_x_work  # noqa: E402
from traceview.project import load_project, save_project  # noqa: E402


@pytest.fixture(scope="module")
def eng(tmp_path_factory):
    rgb, _ = make_sheet()
    p = tmp_path_factory.mktemp("arc") / "s.png"
    cv2.imwrite(str(p), cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
    e = Engine()
    e.open(str(p))
    e.analyze()
    return e


def _cal(**kw):
    c = Calibration(mm_per_px_x=0.1, mm_per_px_y=0.1, arm_length_mm=60.0)
    for k, v in kw.items():
        setattr(c, k, v)
    return c


def test_corrected_path_matches_point_formula(eng):
    b = eng.bands[0]
    cal = _cal()
    cp = corrected_path(b.path, b.line_zero, eng.scale, cal)
    for i in (0, 500, 2000, 4000):
        x, y = b.path[i]
        assert abs(cp[i, 0] - corrected_x_work(float(x), float(y), b.line_zero, eng.scale, cal)) < 1e-6
        assert cp[i, 1] == y                       # only x changes


def test_no_correction_without_arm_or_calibration(eng):
    b = eng.bands[0]
    for cal in (Calibration(), _cal(arm_length_mm=0.0), _cal(mm_per_px_x=0.0)):
        assert not arc_active(cal)
        cp = corrected_path(b.path, b.line_zero, eng.scale, cal)
        assert np.array_equal(cp, b.path.astype(float))


def test_baseline_unchanged_and_large_amplitude_shifts(eng):
    b = eng.bands[0]
    cal = _cal()
    assert abs(corrected_x_work(500.0, b.line_zero, b.line_zero, eng.scale, cal) - 500.0) < 1e-9
    # amplitude 30 mm (300 px at 0.1 mm/px) with a 60 mm arm moves x by 60 - sqrt(60^2-30^2) = 8.04 mm
    d = 500.0 - corrected_x_work(500.0, b.line_zero - 300.0, b.line_zero, eng.scale, cal)
    assert abs(d * 0.1 - 8.0385) < 1e-3
    cal2 = _cal(arc_sign=-1)
    d2 = 500.0 - corrected_x_work(500.0, b.line_zero - 300.0, b.line_zero, eng.scale, cal2)
    assert abs(d2 + d) < 1e-9                      # opposite direction


def test_section_reports_raw_and_corrected_side_by_side(eng):
    b = eng.bands[1]
    b.sections.clear()
    eng.add_section(1, 400, 1200)
    eng.calibration = _cal(seconds_per_mm=0.2)
    r = eng.section_result(1, 0)
    assert r["period_px"] and r["period_arc_px"]
    assert r["period_arc_px"] != r["period_px"]
    assert abs(r["period_arc_s"] - r["period_arc_mm"] * 0.2) < 1e-9
    eng.calibration = Calibration()
    r = eng.section_result(1, 0)
    assert r["period_arc_px"] is None


def test_read_point_snaps_to_trace_and_gives_both_values(eng):
    eng.readouts.clear()
    eng.calibration = _cal(seconds_per_mm=0.2)
    b = eng.bands[0]
    x, y = b.path[3000]
    row = eng.read_point(0, float(x) + 4, float(y) + 3)
    assert row is not None
    rd = eng.readouts[-1]
    on_path = np.any((b.path[:, 0] == rd["x"]) & (b.path[:, 1] == rd["y"]))
    assert on_path                                                    # snapped onto the trace
    assert np.hypot(rd["x"] - (x + 4), rd["y"] - (y + 3)) <= 5.0 + 1e-9  # and not farther than the click's own distance
    assert row["x_arc_mm"] is not None and row["t_arc_s"] is not None
    assert abs(row["x_mm"] - rd["x"] * 0.1) < 1e-6
    assert eng.read_point(0, 5.0, 5.0) is None                          # nothing near
    assert len(eng.readouts) == 1
    assert eng.remove_readout(0) and not eng.readouts
    eng.calibration = Calibration()


def test_readouts_saved_and_arc_in_excel(eng, tmp_path):
    from openpyxl import load_workbook
    eng.readouts.clear()
    eng.calibration = _cal(seconds_per_mm=0.2)
    eng.read_point(0, float(eng.bands[0].path[1000][0]), float(eng.bands[0].path[1000][1]))
    f = str(tmp_path / "w.trproj")
    save_project(eng, f)
    e2 = Engine()
    load_project(e2, f)
    assert len(e2.readouts) == 1
    wb = load_workbook(export_xlsx(e2, str(tmp_path / "x.xlsx")))
    assert "読取点" in wb.sheetnames
    heads = [c.value for c in wb["読取点"][1]]
    assert "円弧補正後のX[mm]" in heads
    assert len(wb["測定結果と画像"]._images) == 3                        # original, trace only, arc overlay
    eng.readouts.clear()
    eng.calibration = Calibration()
