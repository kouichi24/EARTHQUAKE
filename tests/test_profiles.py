import os
import sys

import cv2
import numpy as np
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tests.synth import make_sheet  # noqa: E402
from traceview.engine import Engine  # noqa: E402
from traceview.excel_out import export_xlsx  # noqa: E402
from traceview.measure import arc_corrected_x_mm  # noqa: E402
from traceview.profiles import PROFILES, get  # noqa: E402
from traceview.project import load_project, save_project  # noqa: E402
from traceview.preprocess import prepare  # noqa: E402


def _engine(tmp_path, name, **kw):
    rgb, truth = make_sheet(**kw)
    p = tmp_path / f"{name}.png"
    cv2.imwrite(str(p), cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
    e = Engine()
    e.open(str(p))
    return e, truth


def _recall(band, peaks, tol=6):
    hit = sum(any(v.kind == k and abs(v.x - x) <= tol and abs(v.y - y) <= tol for v in band.vertices)
              for x, y, k in peaks)
    return hit / len(peaks)


def test_all_profiles_exist_and_apply():
    assert [p.key for p in PROFILES] == ["ink_lever", "inverted", "thick", "generic"]
    assert get("ink_lever").apply().drop_marks
    assert get("inverted").apply().invert
    assert not get("generic").apply().drop_marks
    assert get("unknown-key").key == "generic"
    assert all(p.verified for p in PROFILES)


def test_marks_removed_only_when_profile_says_so(tmp_path):
    e, truth = _engine(tmp_path, "ticks", ticks=True)
    gen = e.clone_for_profile(get("generic"))
    gen.ensure_graph()
    lev = e.clone_for_profile(get("ink_lever"))
    lev.ensure_graph()

    def marks_left(eng):
        """How many of the 3 solid marks (16x10 px above each band's reference lines) survive."""
        P = eng.prepared
        n = 0
        for by in (250, 600, 950):
            n += int(P.ink[by - 108:by - 98, 150:167].sum()) > 100
        return n

    assert marks_left(gen) >= 2            # marks are still there without the profile
    assert marks_left(lev) == 0            # removed by the profile


def test_lever_profile_keeps_all_tips_with_marks(tmp_path):
    e, truth = _engine(tmp_path, "ticks2", ticks=True)
    lev = e.clone_for_profile(get("ink_lever"))
    lev.analyze()
    assert len(lev.bands) == 3
    for b, t in zip(lev.bands, truth.values()):
        assert _recall(b, t["peaks"]) >= 0.95


def test_inverted_profile_reads_light_on_dark(tmp_path):
    e, truth = _engine(tmp_path, "inv", invert=True)
    inv = e.clone_for_profile(get("inverted"))
    inv.analyze()
    assert len(inv.bands) == 3
    for b, t in zip(inv.bands, truth.values()):
        assert b.complete
        assert _recall(b, t["peaks"]) >= 0.9


def test_compare_ranks_the_matching_profile_first(tmp_path):
    for kw, expect in (({}, {"ink_lever", "generic"}), ({"invert": True}, {"inverted"})):
        e, _ = _engine(tmp_path, "cmp" + str(len(kw)), **kw)
        rows = []
        for p in PROFILES:
            c = e.clone_for_profile(p)
            c.analyze()
            rows.append((Engine.quality_sort_key(c.quality_summary()), p.key))
        rows.sort()
        assert rows[0][1] in expect, rows


def test_arc_correction_math():
    assert arc_corrected_x_mm(100, 0, 100, 1) == 100
    assert arc_corrected_x_mm(100, 30, 0, 1) == 100               # no arm length -> unchanged
    assert abs(arc_corrected_x_mm(100, 30, 100, 1) - 95.3939) < 1e-3
    assert abs(arc_corrected_x_mm(100, 30, 100, -1) - 104.6061) < 1e-3
    assert arc_corrected_x_mm(100, 120, 100, 1) == 100            # outside the model


def test_profile_saved_and_excel_arc_columns(tmp_path):
    from openpyxl import load_workbook
    e, _ = _engine(tmp_path, "arc")
    c = e.clone_for_profile(get("ink_lever"))
    c.analyze()
    c.calibration.mm_per_px_x = 0.05
    c.calibration.mm_per_px_y = 0.05
    c.calibration.arm_length_mm = 120.0
    c.calibration.seconds_per_mm = 0.5
    f = str(tmp_path / "w.trproj")
    save_project(c, f)
    e2 = Engine()
    load_project(e2, f)
    assert e2.profile_key == "ink_lever"
    assert e2.params.drop_marks
    assert e2.calibration.arm_length_mm == 120.0
    wb = load_workbook(export_xlsx(e2, str(tmp_path / "a.xlsx")))
    heads = [c.value for c in wb["頂点と折返し"][1]]
    assert "円弧補正後のX[mm]" in heads and "補正後Xからの時間[s]" in heads
    # no arm length -> columns absent
    e2.calibration.arm_length_mm = 0.0
    wb = load_workbook(export_xlsx(e2, str(tmp_path / "b.xlsx")))
    assert "円弧補正後のX[mm]" not in [c.value for c in wb["頂点と折返し"][1]]
