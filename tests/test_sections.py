import os
import sys

import cv2
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tests.synth import make_sheet  # noqa: E402
from traceview.engine import Engine, EngineError  # noqa: E402
from traceview.excel_out import export_xlsx  # noqa: E402
from traceview.measure import Calibration  # noqa: E402
from traceview.project import load_project, save_project  # noqa: E402


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    rgb, truth = make_sheet()
    p = tmp_path_factory.mktemp("sec") / "s.png"
    cv2.imwrite(str(p), cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
    e = Engine()
    e.open(str(p))
    e.analyze()
    return e, truth


def test_counts_and_extremes_match_truth(sheet):
    e, truth = sheet
    e.bands[0].sections.clear()
    e.add_section(0, 400, 1200)
    r = e.section_result(0, 0)
    peaks = [p for p in truth[0]["peaks"] if 400 <= p[0] <= 1200 and p[2] == "peak"]
    valleys = [p for p in truth[0]["peaks"] if 400 <= p[0] <= 1200 and p[2] == "valley"]
    assert r["peaks"] == len(peaks)
    assert r["valleys"] == len(valleys)
    assert r["waves"] == len(peaks)
    y0 = e.bands[0].line_zero
    assert abs(r["max_value"] - (y0 - min(p[1] for p in peaks + valleys))) < 4
    assert abs(r["min_value"] - (y0 - max(p[1] for p in peaks + valleys))) < 4
    xs = sorted(p[0] for p in peaks)
    assert abs(r["period_px"] - (xs[-1] - xs[0]) / (len(xs) - 1)) < 3


def test_time_only_when_calibrated(sheet):
    e, _ = sheet
    e.calibration = Calibration()
    r = e.section_result(0, 0)
    assert r["period_s"] is None and r["freq_hz"] is None and r["length_s"] is None
    e.calibration = Calibration(seconds_per_px_x=0.01)
    r = e.section_result(0, 0)
    assert abs(r["period_s"] - r["period_px"] * 0.01) < 1e-9
    assert abs(r["freq_hz"] - 1 / r["period_s"]) < 1e-9
    assert abs(r["length_s"] - r["length_px"] * 0.01) < 1e-9
    # paper speed x mm-per-px gives the same kind of time scale
    e.calibration = Calibration(mm_per_px_x=0.1, seconds_per_mm=0.2)
    assert abs(e.calibration.seconds_per_orig_px() - 0.02) < 1e-12
    e.calibration = Calibration()


def test_small_wiggles_ignored_by_setting(sheet):
    e, _ = sheet
    e.section_settings.min_swing_px = 0.0
    many = e.section_result(0, 0)["waves"]
    e.section_settings.min_swing_px = 60.0
    few = e.section_result(0, 0)["waves"]
    e.section_settings.min_swing_px = 20.0
    assert few <= many


def test_count_modes(sheet):
    e, _ = sheet
    r = {}
    for m in ("peak", "valley", "both"):
        e.section_settings.count_mode = m
        r[m] = e.section_result(0, 0)["waves"]
    e.section_settings.count_mode = "peak"
    assert r["both"] == (r["peak"] + r["valley"]) / 2


def test_short_section_rejected_and_undo(sheet):
    e, _ = sheet
    with pytest.raises(EngineError):
        e.add_section(0, 100, 101)
    n = len(e.bands[0].sections)
    e.add_section(0, 300, 600)
    assert len(e.bands[0].sections) == n + 1
    assert e.undo(0)
    assert len(e.bands[0].sections) == n


def test_sections_saved_and_in_excel(sheet, tmp_path):
    from openpyxl import load_workbook
    e, _ = sheet
    e.calibration = Calibration(seconds_per_px_x=0.01)
    e.section_settings.min_swing_px = 25.0
    f = str(tmp_path / "w.trproj")
    save_project(e, f)
    e2 = Engine()
    load_project(e2, f)
    assert len(e2.bands[0].sections) == len(e.bands[0].sections)
    assert e2.section_settings.min_swing_px == 25.0
    wb = load_workbook(export_xlsx(e2, str(tmp_path / "x.xlsx")))
    ws = wb["区間測定"]
    heads = [c.value for c in next(r for r in ws.iter_rows() if r[0].value == "段")]
    assert "波の数" in heads and "平均周期[s]" in heads
    n_sections = sum(len(b.sections) for b in e2.bands)
    assert n_sections >= 1 and ws.max_row == 8 + n_sections
    e.calibration = Calibration()
    e.section_settings.min_swing_px = 20.0
