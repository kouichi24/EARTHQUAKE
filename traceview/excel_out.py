"""Excel (.xlsx) export with the original image and the overlay (OUT-01, OUT-02)."""
import io
import time

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .measure import arc_corrected_x_mm
from .profiles import get as get_profile

SRC_LABEL = {"auto": "自動取得", "manual": "手動頂点", "retrace": "再追跡"}


class PendingEditsError(Exception):
    pass


def overlay_image(engine, max_w=1800, with_original=True, draw_vertices=True, draw_lines=True):
    img = cv2.cvtColor(engine.work, cv2.COLOR_RGB2BGR) if with_original else \
        np.full(engine.work.shape, 255, np.uint8)
    img = img.copy()
    t = max(2, engine.work.shape[1] // 1600)
    for b in engine.bands:
        c = tuple(int(b.color[i:i + 2], 16) for i in (5, 3, 1))
        if len(b.path) > 1:
            cv2.polylines(img, [b.path.reshape(-1, 1, 2)], False, c, t)
        if draw_vertices:
            for v in b.vertices:
                cv2.circle(img, (int(v.x), int(v.y)), t + 1, (0, 0, 0), -1)
        if draw_lines:
            x0, _y0, x1, _y1 = engine.prepared.paper if engine.prepared else (0, 0, img.shape[1], img.shape[0])
            for y, dash in ((b.line_zero, False), (b.line_top, True), (b.line_bottom, True)):
                cv2.line(img, (x0, int(y)), (x1, int(y)), c, max(1, t // 2))
    s = min(1.0, max_w / img.shape[1])
    if s < 1:
        img = cv2.resize(img, None, fx=s, fy=s, interpolation=cv2.INTER_AREA)
    ok, buf = cv2.imencode(".png", img)
    return buf.tobytes(), img.shape[1], img.shape[0]


def _simplify_idx(path, eps=1.0):
    """Indices of a Douglas-Peucker simplification of the pen path (tips are always kept)."""
    if len(path) <= 2:
        return list(range(len(path)))
    ap = cv2.approxPolyDP(path.reshape(-1, 1, 2).astype(np.int32), eps, False).reshape(-1, 2)
    idx, j = [], 0
    for q in ap:
        while j < len(path) and (path[j, 0] != q[0] or path[j, 1] != q[1]):
            j += 1
        if j >= len(path):
            break
        idx.append(j)
    return idx or list(range(len(path)))


def export_xlsx(engine, out_path, include_all_points=False):
    if engine.has_pending():
        raise PendingEditsError("未適用の頂点修正があります。再追跡を適用するか取り消してから出力してください。")
    inv = 1.0 / engine.scale
    cal = engine.calibration
    unit = cal.unit_y()
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="DDE6F3")
    bold = Font(bold=True)

    # ---- sheet 1
    ws = wb.active
    ws.title = "測定結果と画像"
    ws["A1"] = "TraceView 測定結果"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("元ファイル", engine.src_path), ("ページ", engine.page + 1), ("アプリの版", __version__),
        ("出力日時", time.strftime("%Y-%m-%d %H:%M:%S")),
        ("座標系", "元画像の左上を原点、Xは右向き、Yは下向き（元画像のpx）。振幅 a = y0 − y（0線より上が正）"),
        ("換算条件", "未校正（px表示）" if not cal.calibrated_y else
         f"Y: {cal.mm_per_px_y:.5f} mm/px, X: {cal.mm_per_px_x:.5f} mm/px, 紙送り {cal.seconds_per_mm} s/mm, "
         f"振幅係数 {cal.amplitude_factor} {cal.amplitude_unit} / mm  {cal.note}"),
        ("値の単位", unit),
        ("記録方式の設定", get_profile(engine.profile_key).name + "／" + get_profile(engine.profile_key).verified),
        ("円弧補正", (f"ペンアーム長 {cal.arm_length_mm} mm、向き {cal.arc_sign:+d}（補正後Xの列を出力）"
                    if cal.arm_length_mm > 0 and cal.mm_per_px_x > 0 and cal.mm_per_px_y > 0 else "なし")),
    ]
    for i, (k, v) in enumerate(meta):
        ws.cell(3 + i, 1, k).font = bold
        ws.cell(3 + i, 2, v)
    r0 = 3 + len(meta) + 1
    heads = ["波形番号", "名称", "確認状態", "0線の高さY[px]", "上線の高さY[px]", "下線の高さY[px]", "線の調整状態",
             "最大値", "最大値X[px]", "最大値Y[px]", "最小値", "最小値X[px]", "最小値Y[px]",
             "上振幅", "下振幅", "上下幅", "線基準 上振幅", "線基準 下振幅", "線基準 上下幅",
             "連続頂点間の最大差(参考)", "単位", "末端まで追跡", "要確認の数", "頂点数", "最大・最小の対象外(所属要確認)の頂点数"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(r0, c, h)
        cell.font = bold
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, b in enumerate(engine.bands):
        m = engine.measures(b)
        adj = []
        adj.append("上線:手動" if b.line_top_manual else "上線:自動")
        adj.append("下線:手動" if b.line_bottom_manual else "下線:自動")
        row = [i + 1, b.name, engine.value_status(b), round(b.line_zero * inv, 2), round(b.line_top * inv, 2),
               round(b.line_bottom * inv, 2), " / ".join(adj),
               m.get("max_value"), _r(m.get("max_x")), _r(m.get("max_y")),
               m.get("min_value"), _r(m.get("min_x")), _r(m.get("min_y")),
               m.get("up_amp"), m.get("down_amp"), m.get("span"),
               m.get("line_up"), m.get("line_down"), m.get("line_span"),
               m.get("max_step"), unit, "はい" if b.complete else "いいえ", len(b.issues), len(b.vertices),
               m.get("n_suspect", 0)]
        for c, v in enumerate(row, 1):
            ws.cell(r0 + 1 + i, c, round(v, 3) if isinstance(v, float) else v)
    note_r = r0 + len(engine.bands) + 2
    ws.cell(note_r, 1, "「確認状態」が「確認済み」以外の値は候補です（確定した最大・最小値として扱わないでください）。").font = Font(italic=True)
    png, w, h = overlay_image(engine, 1800)
    img = XLImage(io.BytesIO(png))
    img.width, img.height = w * 0.6, h * 0.6
    ws.add_image(img, f"A{note_r + 2}")
    png2, w2, h2 = overlay_image(engine, 1800, with_original=False)
    img2 = XLImage(io.BytesIO(png2))
    img2.width, img2.height = w2 * 0.6, h2 * 0.6
    ws.add_image(img2, f"A{note_r + 2 + int(h * 0.6 / 20) + 3}")
    for c in range(1, len(heads) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 22

    # ---- sheet 2: reconstructed coordinates
    ws2 = wb.create_sheet("復元座標")
    arc_on = cal.arm_length_mm > 0 and cal.mm_per_px_x > 0 and cal.mm_per_px_y > 0
    h2s = ["段", "復元順", "X[px]", "Y[px]", f"振幅({unit})", "区分"] + (["円弧補正後のX[mm]"] if arc_on else []) \
        + (["補正後Xからの時間[s]"] if arc_on and cal.seconds_per_mm > 0 else [])
    for c, h in enumerate(h2s, 1):
        cell = ws2.cell(1, c, h)
        cell.font = bold
        cell.fill = hdr_fill
    r = 2
    for b in engine.bands:
        idxs = range(len(b.path)) if include_all_points else sorted(set(_simplify_idx(b.path)) | {v.idx for v in b.vertices})
        vmap = {v.idx: v for v in b.vertices}
        for i in idxs:
            x, y = b.path[i]
            kind = "自動取得"
            if b.is_inferred_index(i):
                kind = "推定接続"
            v = vmap.get(i)
            if v is not None and v.source != "auto":
                kind = SRC_LABEL[v.source]
            row2 = [b.name, int(i) + 1, round(float(x) * inv, 2), round(float(y) * inv, 2),
                    round(cal.to_y_units((b.line_zero - float(y)) * inv), 3), kind]
            if arc_on:
                row2 += _arc_cols(cal, b, float(x), float(y), inv)
            ws2.append(row2)
            r += 1
    ws2.freeze_panes = "A2"

    # ---- sheet 3: vertices
    ws3 = wb.create_sheet("頂点と折返し")
    h3 = ["段", "頂点番号", "X[px]", "Y[px]", "山/谷", f"振幅({unit})", "往路/復路", "折返しの種類",
          "対応する復元順", "由来", "要確認事項"] + (["円弧補正後のX[mm]"] if arc_on else []) \
        + (["補正後Xからの時間[s]"] if arc_on and cal.seconds_per_mm > 0 else [])
    for c, h in enumerate(h3, 1):
        cell = ws3.cell(1, c, h)
        cell.font = bold
        cell.fill = hdr_fill
    for b in engine.bands:
        for n, v in enumerate(b.vertices, 1):
            note = v.note
            if b.is_inferred_index(v.idx):
                note = (note + " 推定接続の区間上").strip()
            row3 = [b.name, n, round(v.x * inv, 2), round(v.y * inv, 2), "山" if v.kind == "peak" else "谷",
                    round(cal.to_y_units((b.line_zero - v.y) * inv), 3),
                    "往路で到達" + (f"・復路で再通過{v.revisit}回" if v.revisit else ""),
                    "先端で反転(行き止まり)" if v.turn == "end" else "曲がり",
                    v.idx + 1, SRC_LABEL.get(v.source, v.source), note]
            if arc_on:
                row3 += _arc_cols(cal, b, v.x, v.y, inv)
            ws3.append(row3)
    ws3.freeze_panes = "A2"

    # ---- sheet 4: issues
    ws4 = wb.create_sheet("要確認")
    for c, h in enumerate(["段", "種類", "X[px]", "Y[px]", "内容"], 1):
        cell = ws4.cell(1, c, h)
        cell.font = bold
        cell.fill = hdr_fill
    kinds = {"ambiguous": "所属の曖昧さ", "untracked_end": "末端の未追跡", "inferred": "推定接続",
             "excursion": "他の段への逸脱", "no_origin": "起点未検出", "retrace_failed": "再追跡の失敗"}
    for b in engine.bands:
        for it in b.issues:
            ws4.append([b.name, kinds.get(it.kind, it.kind), round(it.x * inv, 1), round(it.y * inv, 1), it.msg])
    for w_, col in zip((12, 20, 12, 12, 60), "ABCDE"):
        ws4.column_dimensions[col].width = w_
    wb.save(out_path)
    return out_path


def _r(v):
    return None if v is None else round(v, 2)


def _arc_cols(cal, band, x, y, inv):
    """[arc-corrected x in mm, time in s from the band origin] for one point (original px = work * inv)."""
    amp_mm = (band.line_zero - y) * inv * cal.mm_per_px_y
    x_mm = arc_corrected_x_mm(x * inv * cal.mm_per_px_x, amp_mm, cal.arm_length_mm, cal.arc_sign)
    out = [round(x_mm, 3)]
    if cal.seconds_per_mm > 0:
        x0_mm = band.origin[0] * inv * cal.mm_per_px_x
        out.append(round((x_mm - x0_mm) * cal.seconds_per_mm, 3))
    return out
